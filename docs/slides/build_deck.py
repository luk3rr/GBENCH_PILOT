#!/usr/bin/env python3
"""
Filename: build_deck.py
Created on: September 27, 2026
Author: Lucas Araújo <araujolucas@dcc.ufmg.br>

Generates google_benchmark.pptx from results/summary.json.
Requires: pip install python-pptx
"""

import json
import os
import re

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
DATA = json.load(open(os.path.join(ROOT, "results", "summary.json")))

# Palette ---------------------------------------------------------------------
INK = "12161C"  # dark background
PANEL = "1F2630"  # code panels
ORANGE = "FF6B2C"  # accent: speed / attention
TEAL = "16A394"
PURPLE = "7B5CC4"
GRAY = "8A96A3"
TEXT = "1B222A"
MUTED = "5B6772"
TINT = "F0F3F6"
WHITE = "FFFFFF"
CODE_FG = "E6EAF0"
CODE_COMMENT = "7F8C99"
CODE_KW = "5CC8BC"
CODE_API = "FF9A66"
CODE_STR = "B5D67A"

SANS = "Arial"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW = 13.333


def rgb(h):
    return RGBColor.from_string(h)


# Primitives ------------------------------------------------------------------
def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, line=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def text(slide, x, y, w, h, content, size=16, color=TEXT, bold=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=None, margin=0):
    """content: str | list of paragraphs; paragraph: str | list of runs (text, {opts})"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))

    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        runs = para if isinstance(para, list) else [(para, {})]
        for rtext, opts in runs:
            r = p.add_run()
            r.text = rtext
            f = r.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.color.rgb = rgb(opts.get("color", color))
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=TEXT, spacing=10):
    tb = text(slide, x, y, w, h, items, size=size, color=color, spacing=spacing)
    for p in tb.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.28)))
        pPr.set("indent", str(-Inches(0.28)))
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "•")
    return tb


def title(slide, t, sub=None, dark=False):
    text(slide, 0.6, 0.45, 12.1, 0.8, t, size=34, bold=True, color=WHITE if dark else TEXT)
    if sub:
        text(slide, 0.6, 1.2, 12.1, 0.45, sub, size=16, color=GRAY if dark else MUTED)


def badge(slide, x, y, label, d=0.42, fill=ORANGE, color=WHITE, size=14):
    c = box(slide, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = SANS
    r.font.color.rgb = rgb(color)
    return c


KEYWORDS = r"\b(for|auto|static|void|const|bool|class|public|protected|override|return|if|while|std::size_t|constexpr|template|typename|INCLUDE|FetchContent_Declare|FetchContent_MakeAvailable|TARGET_LINK_LIBRARIES|SET)\b"
API = r"(benchmark::\w+|BENCHMARK\w*|state\.\w+|->\w+|\bstate\b)"


def highlight(line):
    """Tiny C++/CMake tokenizer: returns [(text, color)]"""
    stripped = line.lstrip()
    if stripped.startswith("//") or stripped.startswith("#"):
        return [(line, CODE_COMMENT)]
    idx = line.find("//")
    tail = []
    if idx >= 0:
        line, tail = line[:idx], [(line[idx:], CODE_COMMENT)]
    pattern = re.compile(f"{API}|{KEYWORDS}|(\"[^\"]*\")")
    out, pos = [], 0
    for m in pattern.finditer(line):
        if m.start() > pos:
            out.append((line[pos:m.start()], CODE_FG))
        if m.group(1):
            col = CODE_API
        elif m.group(2):
            col = CODE_KW
        else:
            col = CODE_STR
        out.append((m.group(0), col))
        pos = m.end()
    if pos < len(line):
        out.append((line[pos:], CODE_FG))
    return out + tail


def code(slide, x, y, w, h, src, size=12, label=None, plain=False, pad=0.25):
    box(slide, x, y, w, h, PANEL, radius=0.04)
    top = y + pad
    if label:
        text(slide, x + pad, y + 0.14, w - 2 * pad, 0.3, label, size=10, color=GRAY, bold=True, font=SANS)
        top = y + 0.45
    paras = []
    for line in src.strip("\n").split("\n"):
        if plain:
            paras.append([(line if line else " ", {"color": CODE_FG})])
        else:
            runs = highlight(line) if line else [(" ", CODE_FG)]
            paras.append([(t, {"color": c}) for t, c in runs])
    tb = text(slide, x + pad, top, w - 2 * pad, h - (top - y) - 0.15, paras, size=size, font=MONO)
    for p in tb.text_frame.paragraphs:
        p.line_spacing = 1.08
    return tb


def card(slide, x, y, w, h, head, body, head_color=TEXT, fill=TINT, size=13, head_size=15):
    box(slide, x, y, w, h, fill, radius=0.08)
    text(slide, x + 0.25, y + 0.2, w - 0.5, 0.4, head, size=head_size, bold=True, color=head_color)
    text(slide, x + 0.25, y + 0.62, w - 0.5, h - 0.75, body, size=size, color=MUTED)


def stat(slide, x, y, w, value, label, color=ORANGE, vsize=54, lsize=14, lcolor=MUTED):
    text(slide, x, y, w, 0.95, value, size=vsize, bold=True, color=color)
    text(slide, x, y + 0.98, w, 0.6, label, size=lsize, color=lcolor)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def br(v, nd=1):
    return f"{v:.{nd}f}".replace(".", ",")


def fmt_n(n):
    return f"{n:,}".replace(",", ".")


def chart(slide, x, y, w, h, kind, cats, series, colors, log=False, legend=True,
          number_format="0", val_title=None, labels=False, font_size=12):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.size = Pt(font_size)
    ch.font.name = SANS
    ch.font.color.rgb = rgb(MUTED)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(font_size)

    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("DDE2E7")
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.number_format = number_format
    va.tick_labels.number_format_is_linked = False
    if val_title:
        va.has_title = True
        va.axis_title.text_frame.text = val_title
        r = va.axis_title.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(font_size)
        r.font.bold = False
        r.font.color.rgb = rgb(MUTED)
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb("B8C0C8")
    ca.has_major_gridlines = False

    if log:
        scaling = va._element.find(qn("c:scaling"))
        lb = etree.SubElement(scaling, qn("c:logBase"))
        lb.set("val", "10")
        scaling.remove(lb)
        scaling.insert(0, lb)  # logBase must come first inside c:scaling

    plot = ch.plots[0]
    for s, col in zip(plot.series, colors):
        fmt = s.format
        if kind in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
            fmt.line.color.rgb = rgb(col)
            fmt.line.width = Pt(3)
            s.smooth = False
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = 7
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = rgb(col)
            s.marker.format.line.color.rgb = rgb(col)
        else:
            fmt.fill.solid()
            fmt.fill.fore_color.rgb = rgb(col)
    if kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
        plot.gap_width = 60
        plot.overlap = -10
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = number_format
        dl.number_format_is_linked = False
        dl.font.size = Pt(font_size - 1)
        dl.font.color.rgb = rgb(TEXT)
    return ch


D = DATA
N_LABELS = [fmt_n(n) for n in D["N"]]
BIGO = D["bigo"]
BIGO_V0 = D["bigo_v0"]


def rms(d, name):
    return round(d[f"{name}_RMS"][2] * 100)


# 1. Title ---------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, INK)
text(s, 0.8, 1.35, 6.5, 0.4, "C++  ·  PERFORMANCE  ·  MICROBENCHMARKS", size=13, bold=True, color=ORANGE)
text(s, 0.8, 1.85, 6.4, 1.9, ["Google", "Benchmark"], size=54, bold=True, color=WHITE)
text(s, 0.8, 3.95, 6.2, 1.1, "Medindo desempenho de código C++ sem se enganar", size=22, color="C9D1D9")
text(s, 0.8, 5.35, 6, 0.4, "Lucas Araújo", size=16, bold=True, color=WHITE)
text(s, 0.8, 5.75, 6, 0.4, "Setembro de 2026", size=14, color=GRAY)
code(s, 7.55, 1.6, 5.2, 3.9, f"""
$ ./benchmarks
Benchmark               Time
-----------------------------
BM_LinearSearch/16    {D['linear'][0]:.1f} ns
BM_LinearSearch/1024   {D['linear'][3]:.0f} ns
BM_LinearSearch/65536 {D['linear'][6]:.0f} ns
BM_LinearSearch_BigO  0.17 N

BM_BinarySearch/65536 {D['binary'][6]:.1f} ns
BM_BinarySearch_BigO  4.00 lgN
""", size=14, plain=True)
notes(s, "Apresentação sobre a biblioteca Google Benchmark: o que é, como usar, "
         "e um piloto prático medindo diferentes implementações de busca num projeto C++ próprio. "
         "Os números no terminal à direita são reais, da minha máquina.")

# 2. Why --------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Por que não só std::chrono?", "O jeito ingênuo de medir tempo tem armadilhas que não aparecem no número final")
code(s, 0.6, 1.95, 5.6, 4.9, """
auto t0 = std::chrono::steady_clock::now();

for (int i = 0; i < 1000; ++i)
    LinearContains(values, 42);

auto t1 = std::chrono::steady_clock::now();

std::cout << (t1 - t0) / 1000 << "\\n";
// Mediu o quê, exatamente?
""", size=13, label="MEDIÇÃO MANUAL")
problems = [
    ("O compilador pode apagar o que você mede",
     "O resultado não é usado; com -O3 o laço inteiro some."),
    ("Quantas iterações? 1000 é chute",
     "Pouco demais vira ruído; demais desperdiça tempo."),
    ("Sem aquecimento",
     "Cache, TLB e frequência da CPU ainda estão \"frios\"."),
    ("Sem estatística",
     "Uma execução só. Ficou mais rápido mesmo ou foi sorte?"),
]
for i, (h, b) in enumerate(problems):
    yy = 2.0 + i * 1.22
    badge(s, 6.75, yy, str(i + 1))
    text(s, 7.4, yy - 0.04, 5.4, 0.4, h, size=17, bold=True)
    text(s, 7.4, yy + 0.4, 5.4, 0.6, b, size=14, color=MUTED)
notes(s, "Todo mundo já fez isso. O problema é que esse código tem pelo menos quatro fontes de erro. "
         "A mais grave é a primeira: como o retorno não é usado, o otimizador pode remover a chamada inteira, "
         "e a gente vai mostrar isso acontecendo no piloto.")

# 3. What is ---------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "O que é o Google Benchmark")
text(s, 0.6, 1.45, 4.3, 3.2, [
    "Biblioteca open source do Google (licença Apache 2.0) para microbenchmarks em C++.",
    "Você escreve só o trecho a medir. A biblioteca cuida do resto: quantas vezes rodar, "
    "como agregar e como exportar.",
], size=17, spacing=14, color=TEXT)
text(s, 0.6, 4.75, 4.3, 1.5, [[("github.com/google/benchmark", {"bold": True, "color": ORANGE})],
                              "Versão usada no piloto: v1.9.5"], size=14, color=MUTED, spacing=6)
feats = [
    ("Iterações automáticas", "Roda até o tempo estabilizar; você não escolhe o N de repetições."),
    ("Parâmetros e ranges", "Um benchmark, vários tamanhos: Arg, Range, DenseRange."),
    ("Big-O automático", "Ajusta O(1), O(N), O(N log N)… aos dados e reporta o erro."),
    ("Fixtures e templates", "Setup compartilhado e o mesmo teste para vários tipos."),
    ("Counters", "items/s, bytes/s ou qualquer métrica sua como coluna."),
    ("JSON/CSV + compare.py", "Exporta resultados e compara A/B com teste estatístico."),
]
for i, (h, b) in enumerate(feats):
    col, row = i % 2, i // 2
    card(s, 5.35 + col * 3.8, 1.45 + row * 1.85, 3.55, 1.6, h, b, size=13)
notes(s, "A ideia central: você escreve o corpo do laço e a biblioteca decide quantas vezes executar. "
         "Os seis recursos à direita são os que eu usei no piloto.")

# 4. Setup ------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Setup com CMake", "Sem instalar nada no sistema: o CMake baixa e compila junto com o projeto")
code(s, 0.6, 1.95, 7.3, 4.95, """
INCLUDE(FetchContent)
SET(BENCHMARK_ENABLE_TESTING OFF CACHE BOOL "" FORCE)

FetchContent_Declare(
    googlebenchmark
    GIT_REPOSITORY https://github.com/google/benchmark.git
    GIT_TAG v1.9.5
    GIT_SHALLOW TRUE
)
FetchContent_MakeAvailable(googlebenchmark)

ADD_EXECUTABLE(benchmarks ${BENCHMARKS})
TARGET_LINK_LIBRARIES(benchmarks searchlib
    benchmark::benchmark benchmark::benchmark_main)
""", size=12, label="CMakeLists.txt")
steps = [
    ("FetchContent", "Clona a tag v1.9.5 e compila como subprojeto."),
    ("benchmark_main", "Já fornece o main(); o arquivo só tem os benchmarks."),
    ("Sempre em Release", "Medir -O0 não diz nada sobre produção."),
]
for i, (h, b) in enumerate(steps):
    yy = 2.0 + i * 1.2
    badge(s, 8.35, yy, str(i + 1))
    text(s, 9.0, yy - 0.04, 3.8, 0.4, h, size=17, bold=True)
    text(s, 9.0, yy + 0.4, 3.8, 0.7, b, size=14, color=MUTED)
box(s, 8.35, 5.65, 4.4, 1.25, TINT, radius=0.08)
text(s, 8.6, 5.8, 4.0, 1.0, [[("Alternativas: ", {"bold": True, "color": TEXT})],
                            "apt install libbenchmark-dev · vcpkg · conan"], size=13, color=MUTED, spacing=4)
notes(s, "Esse é o trecho real do CMakeLists do piloto. O detalhe importante é o benchmark_main, "
         "que evita escrever o main. E sempre medir em Release: no meu script ./run -B isso é forçado.")

# 5. Anatomy ----------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Anatomia de um benchmark")
code(s, 0.6, 1.4, 7.4, 5.5, """
static void BM_LinearSearch(benchmark::State& state)
{
    const auto values  = MakeSortedEvens(state.range(0));
    const auto queries = MakeQueries(state.range(0), kN);

    std::size_t i = 0;
    for (auto _ : state)
    {
        bool found = LinearContains(values, queries[i]);
        benchmark::DoNotOptimize(found);
        i = (i + 1) % queries.size();
    }

    state.SetItemsProcessed(state.iterations());
    state.SetComplexityN(state.range(0));
}
BENCHMARK(BM_LinearSearch)
    ->RangeMultiplier(4)->Range(16, 1 << 16)
    ->Complexity();
""", size=12.5)
parts = [
    ("Setup fora do laço", "Montar os dados não entra na medição."),
    ("for (auto _ : state)", "Só o corpo é cronometrado. A lib escolhe quantas vezes rodar."),
    ("DoNotOptimize", "Obriga o compilador a produzir o valor, então ele não pode apagar a chamada."),
    ("Counters e Big-O", "Vazão (items/s) e o N usado no ajuste de complexidade."),
    ("Registro", "N = 16, 64, …, 65.536, e no final o Big-O ajustado."),
]
for i, (h, b) in enumerate(parts):
    yy = 1.45 + i * 1.1
    badge(s, 8.4, yy, str(i + 1))
    text(s, 9.05, yy - 0.04, 3.9, 0.4, h, size=16, bold=True)
    text(s, 9.05, yy + 0.36, 3.9, 0.7, b, size=13, color=MUTED)
notes(s, "Esse é o benchmark real do piloto. O range(0) é o parâmetro N que vem do Range. "
         "O laço for (auto _ : state) é o coração: a biblioteca roda até ter uma medição estável, "
         "de algumas centenas de milissegundos por padrão.")

# 6. Output -----------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Lendo a saída", "Execução real no Ryzen 5 5600X, filtrada para a busca linear (cabeçalho resumido)")
code(s, 0.6, 1.95, 12.15, 2.75, """
Run on (12 X 4245.21 MHz CPU s)    CPU Caches: L1 32 KiB · L2 512 KiB · L3 32 MiB
--------------------------------------------------------------------------------------------
Benchmark                  Time             CPU   Iterations UserCounters...
--------------------------------------------------------------------------------------------
BM_LinearSearch/16      8.29 ns         8.29 ns     84806745 items_per_second=120.589M/s
BM_LinearSearch/1024     187 ns          186 ns      3917884 items_per_second=5.36194M/s
BM_LinearSearch/65536  10877 ns        10876 ns        61383 items_per_second=91.9456k/s
BM_LinearSearch_BigO    0.17 N          0.17 N
BM_LinearSearch_RMS        1 %             1 %
""", size=12, plain=True)
cols = [
    ("Time", "Tempo de relógio por iteração."),
    ("CPU", "Tempo de CPU da thread. Se difere de Time, houve espera."),
    ("Iterations", "Escolhido automaticamente: 84 mi para N=16, 61 mil para N=65.536."),
    ("BigO / RMS", "Coeficiente ajustado (0,17·N) e o erro do ajuste."),
]
for i, (h, b) in enumerate(cols):
    card(s, 0.6 + i * 3.1, 5.0, 2.85, 1.7, h, b, head_color=ORANGE, size=13)
notes(s, "Repare nas iterações: a biblioteca rodou 84 milhões de vezes o caso pequeno e 61 mil o grande. "
         "Tudo isso para acumular tempo suficiente para uma medição estável. "
         "O RMS de 1% diz que o modelo O(N) explica muito bem os dados.")

# 7. Parametrize ------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Parametrizando", "Três formas de reaproveitar um benchmark")
cols = [
    ("Argumentos", TEAL, """
->Arg(1024)
->Args({1024, 8})
->Range(16, 1 << 16)
->RangeMultiplier(4)
->DenseRange(0, 64, 8)

// no corpo:
state.range(0)
state.range(1)
"""),
    ("Templates", ORANGE, """
template <typename C>
static void BM_Accumulate(
    benchmark::State& state)
{ ... }

BENCHMARK_TEMPLATE(BM_Accumulate,
    std::vector<int32_t>);
BENCHMARK_TEMPLATE(BM_Accumulate,
    std::list<int32_t>);
"""),
    ("Fixtures", PURPLE, """
class HashFix
  : public benchmark::Fixture {
  void SetUp(const State&)
      override { ... }
};

BENCHMARK_DEFINE_F(HashFix,
    BM_Lookup)(State& st)
{ ... }
BENCHMARK_REGISTER_F(HashFix,
    BM_Lookup)->Range(16, 1<<16);
"""),
]
for i, (h, col, src) in enumerate(cols):
    x = 0.6 + i * 4.1
    badge(s, x, 1.95, "", d=0.26, fill=col)
    text(s, x + 0.4, 1.88, 3.4, 0.4, h, size=18, bold=True)
    code(s, x, 2.5, 3.85, 4.1, src, size=12)
notes(s, "Argumentos: um benchmark vira uma família, um por tamanho. Templates: mesmo código com tipos diferentes; "
         "usei para vector contra list. Fixtures: quando o setup é pesado ou compartilhado, como montar o unordered_set.")

# 8. Beyond time ------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Medindo além do tempo")
rows = [
    ("Complexidade", "->Complexity()   ou   ->Complexity(benchmark::oNLogN)",
     "Com state.SetComplexityN(n), gera as linhas _BigO e _RMS."),
    ("Vazão", "state.SetItemsProcessed(...)   state.SetBytesProcessed(...)",
     "Vira items_per_second e bytes_per_second na saída."),
    ("Counters próprios", "state.counters[\"hit_rate\"] = benchmark::Counter(v);",
     "Qualquer métrica do domínio, com flags de taxa e média."),
    ("Controle do timer", "state.PauseTiming();  /* prepara */  state.ResumeTiming();",
     "Tira preparo da medição, mas tem custo próprio: use com cuidado."),
    ("Threads e unidade", "->Threads(4)   ->UseRealTime()   ->Unit(benchmark::kMicrosecond)",
     "Mede código concorrente e ajusta a unidade exibida."),
]
for i, (h, c, b) in enumerate(rows):
    yy = 1.45 + i * 1.15
    box(s, 0.6, yy, 12.15, 1.0, TINT, radius=0.1)
    text(s, 0.9, yy + 0.1, 3.0, 0.8, h, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.9, yy + 0.12, 8.6, 0.4, c, size=13, font=MONO, color=ORANGE, bold=True)
    text(s, 3.9, yy + 0.52, 8.6, 0.4, b, size=13, color=MUTED)
notes(s, "O ajuste de complexidade é um dos recursos mais interessantes: a biblioteca faz uma regressão "
         "e diz qual Big-O melhor explica os dados. O RMS é o erro desse ajuste; vamos ver que ele serve até para detectar benchmark errado.")

# 9. CLI --------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Linha de comando e ambiente")
flags = [
    ("--benchmark_filter=Search", "regex: roda só o que casar"),
    ("--benchmark_repetitions=10", "repete tudo, gera mean/median/stddev/cv"),
    ("--benchmark_report_aggregates_only", "mostra só os agregados"),
    ("--benchmark_min_time=2s", "tempo mínimo por benchmark (ou 1000x)"),
    ("--benchmark_min_warmup_time=0.5", "aquece antes de medir"),
    ("--benchmark_out=run.json", "salva JSON/CSV para comparar depois"),
    ("--benchmark_list_tests", "lista os benchmarks registrados"),
]
for i, (f, d) in enumerate(flags):
    yy = 1.5 + i * 0.74
    text(s, 0.6, yy, 4.1, 0.4, f, size=13, font=MONO, bold=True, color=TEXT)
    text(s, 4.75, yy + 0.02, 3.7, 0.6, d, size=13, color=MUTED)
box(s, 8.75, 1.45, 4.0, 5.45, INK, radius=0.05)
text(s, 9.05, 1.7, 3.5, 0.4, "Ambiente de medição", size=17, bold=True, color=WHITE)
bullets(s, 9.05, 2.3, 3.5, 4.5, [
    "Build Release (-O3)",
    "Fixar core: taskset -c 2",
    "Governor em performance",
    "Máquina ociosa: nada de compilar ou abrir o navegador enquanto mede",
    "Olhar mediana e CV, não uma execução",
], size=14, color="D5DCE3", spacing=12)
notes(s, "As flags que eu mais usei no piloto. E o ambiente importa tanto quanto o código: "
         "o CV (coeficiente de variação) diz se a medição está confiável. No piloto ficou abaixo de 1% na maioria dos casos.")

# 10. Section: pilot --------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, INK)
text(s, 0.8, 1.6, 6, 0.4, "PILOTO", size=14, bold=True, color=ORANGE)
text(s, 0.8, 2.1, 11.5, 1.2, "Busca em vetor ordenado", size=48, bold=True, color=WHITE)
text(s, 0.8, 3.2, 11, 0.6, "Uma biblioteca pequena, quatro implementações e três pegadinhas",
     size=20, color="C9D1D9")
for i, (v, l) in enumerate([("4", "implementações\nde busca"), ("7", "tamanhos\n16 → 65.536"),
                            ("4,4×", "ganho da versão\nbranchless")]):
    stat(s, 0.8 + i * 3.6, 4.55, 3.3, v, l, color=ORANGE, vsize=48, lsize=15, lcolor=GRAY)
notes(s, "Agora o piloto. Montei um projeto a partir do meu template C++ (CPP_SKELETON) "
         "e usei a biblioteca para tomar decisões sobre uma implementação de busca.")

# 11. Project ---------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "O projeto", "GBENCH_PILOT: baseado no CPP_SKELETON, com doctest para correção e Google Benchmark para desempenho")
code(s, 0.6, 1.95, 4.9, 3.7, """
GBENCH_PILOT/
├── include/search.h
├── src/search.cc
├── test/
│   ├── unit/main_doctest.cc
│   └── benchmarks/
│       ├── bm_search.cc
│       └── bm_pitfalls.cc
├── results/*.json
└── run            # ./run -B
""", size=13, plain=True)
impls = [
    ("LinearContains", "O(N)", "Varre o vetor inteiro", GRAY),
    ("BinaryContains", "O(log N)", "Binária clássica com if", TEAL),
    ("BranchlessBinaryContains", "O(log N)", "Binária sem desvio (cmov)", ORANGE),
    ("std::unordered_set", "O(1)", "Hash, via fixture", PURPLE),
]
for i, (n, o, d, col) in enumerate(impls):
    yy = 1.95 + i * 0.95
    box(s, 5.85, yy, 6.9, 0.8, TINT, radius=0.12)
    badge(s, 6.05, yy + 0.24, "", d=0.32, fill=col)
    text(s, 6.55, yy + 0.08, 4.1, 0.35, n, size=15, bold=True, font=MONO)
    text(s, 6.55, yy + 0.43, 4.1, 0.35, d, size=13, color=MUTED)
    text(s, 10.6, yy + 0.1, 1.95, 0.6, o, size=17, bold=True, color=col, align=PP_ALIGN.RIGHT,
         anchor=MSO_ANCHOR.MIDDLE)
text(s, 5.85, 5.9, 6.9, 1.0, [[("Ambiente: ", {"bold": True, "color": TEXT}),
                              ("Ryzen 5 5600X · g++ 15.2 -O3 · taskset -c 2 · governor performance · "
                               "5 a 10 repetições, mediana · 65.536 consultas aleatórias", {})]],
     size=13, color=MUTED)
notes(s, "Cada busca responde se um valor está no vetor. Os testes unitários (2018 asserções) garantem que as "
         "quatro concordam antes de medir. Não adianta ser rápido e errado.")

# 12. Result: complexity ----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "A complexidade aparece nos números", "Tempo por consulta (ns, escala log) × tamanho do vetor")
chart(s, 0.45, 1.8, 8.3, 5.3, XL_CHART_TYPE.LINE_MARKERS, N_LABELS,
      [("Linear", D["linear"]), ("Binária", D["binary"]), ("unordered_set", D["hash"])],
      [GRAY, TEAL, PURPLE], log=True, number_format="0")
text(s, 9.1, 1.85, 3.7, 0.4, "Big-O que a lib ajustou", size=16, bold=True)
fits = [
    ("Linear", f"{BIGO['BM_LinearSearch_BigO'][0]:.2f}·N", rms(BIGO, "BM_LinearSearch"), GRAY),
    ("Binária", f"{BIGO['BM_BinarySearch_BigO'][0]:.2f}·lgN", rms(BIGO, "BM_BinarySearch"), TEAL),
    ("Hash", f"{BIGO['HashFixture/BM_HashLookup_BigO'][0]:.2f}·(1)", rms(BIGO, "HashFixture/BM_HashLookup"), PURPLE),
]
for i, (n, f, r, col) in enumerate(fits):
    yy = 2.4 + i * 0.72
    text(s, 9.1, yy, 1.3, 0.4, n, size=14, color=MUTED)
    text(s, 10.35, yy, 1.6, 0.4, f, size=14, bold=True, font=MONO, color=col)
    text(s, 11.95, yy, 0.85, 0.4, f"{r}%", size=13, color=MUTED, align=PP_ALIGN.RIGHT)
box(s, 9.1, 4.75, 3.65, 2.1, TINT, radius=0.08)
text(s, 9.35, 4.92, 3.2, 1.9, [
    [("Constantes importam. ", {"bold": True, "color": TEXT}),
     (f"Em N=16 a linear ({br(D['linear'][0])} ns) vence a binária ({br(D['binary'][0])} ns): "
      "varrer 16 inteiros contíguos é barato.", {})]], size=13, color=MUTED)
notes(s, f"Linear cresce 4x a cada passo, a binária quase não se mexe e o hash é praticamente constante. "
         f"O % à direita é o RMS do ajuste. Mas atenção ao começo do gráfico: para vetores pequenos "
         f"a busca linear ganha da binária. Big-O não diz qual é mais rápido no seu N.")

# 13. Pitfall 1 -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Pegadinha 1: o compilador apagou o benchmark", "Mesmo laço, soma de i² para i < 65.536")
code(s, 0.6, 1.95, 6.3, 2.35, """
for (auto _ : state) {
    int64_t sum = 0;
    for (int64_t i = 0; i < n; ++i)
        sum += i * i;          // nunca usado
}
""", size=13, label="SEM DoNotOptimize")
code(s, 0.6, 4.5, 6.3, 2.4, """
for (auto _ : state) {
    int64_t sum = 0;
    for (int64_t i = 0; i < n; ++i) {
        sum += i * i;
        benchmark::DoNotOptimize(sum);
    }
}
""", size=13, label="COM DoNotOptimize")
stat(s, 7.5, 2.0, 5.2, "0,000 ns", "sem DoNotOptimize: -O3 viu que sum não é usado e removeu o laço",
     color=GRAY, vsize=54)
stat(s, 7.5, 4.5, 5.2, f"{br(D['sum_yes'] / 1000)} µs",
     "com DoNotOptimize: o trabalho real, 65.536 multiplicações e somas", color=ORANGE, vsize=54)
notes(s, "Esse é o erro mais clássico. O número de cima parece incrível e é falso: o compilador provou que o laço "
         "não tem efeito e o apagou. Detalhe da v1.9.5: DoNotOptimize com referência const foi descontinuado, "
         "e o compilador deu warning no meu piloto até eu trocar para variável não-const.")

# 14. Pitfall 2 -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Pegadinha 2: o hardware decorou o teste", "Busca binária, 1ª versão (N consultas em ciclo) × versão corrigida (65.536 consultas)")
chart(s, 0.45, 1.8, 7.9, 5.3, XL_CHART_TYPE.LINE_MARKERS, N_LABELS,
      [("1ª versão: N consultas", D["binary_predictable"]), ("Corrigida: 65.536 consultas", D["binary"])],
      [GRAY, TEAL], log=True, number_format="0")
ratio = D["binary"][3] / D["binary_predictable"][3]
stat(s, 8.75, 1.9, 4.0, f"{br(ratio)}× otimista", "em N=1.024: 9,7 ns medidos, 37,5 ns reais",
     color=ORANGE, vsize=40)
stat(s, 8.75, 3.55, 4.0, f"RMS {rms(BIGO_V0, 'BM_BinarySearch')}% → {rms(BIGO, 'BM_BinarySearch')}%",
     "erro do ajuste O(log N)", color=TEXT, vsize=32)
box(s, 8.75, 5.1, 4.0, 1.8, TINT, radius=0.08)
text(s, 9.0, 5.25, 3.55, 1.6, [[("Lição: ", {"bold": True, "color": TEXT}),
                               ("com poucas consultas repetidas, o preditor de desvios da CPU aprende o padrão. "
                                "Big-O que não encaixa é sinal de benchmark errado.", {})]],
     size=13, color=MUTED)
notes(s, "Esse eu descobri durante o piloto. O degrau entre 1024 e 4096 não fazia sentido para O(log N), e o RMS "
         "de 55% denunciou. A causa: com só N consultas em ciclo, o preditor de desvios da Zen 3 memorizava a "
         "sequência. Usando sempre 65.536 consultas, a curva ficou logarítmica e o RMS caiu para 7%.")

# 15. Pitfall 3 -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Pegadinha 3: mesmo O(N), custo 10× diferente", "std::accumulate em vector × list: ns por elemento (BENCHMARK_TEMPLATE)")
per_vec = [round(t / n, 3) for t, n in zip(D["vec"], D["accN"])]
per_list = [round(t / n, 3) for t, n in zip(D["list"], D["accN"])]
chart(s, 0.45, 1.8, 7.9, 5.3, XL_CHART_TYPE.COLUMN_CLUSTERED, [fmt_n(n) for n in D["accN"]],
      [("vector", per_vec), ("list", per_list)], [TEAL, ORANGE], number_format="0.00", labels=True)
vec_bw = D["accN"][2] * 4 / D["vec"][2]  # bytes/ns == GB/s
list_bw = D["accN"][2] * 4 / D["list"][2]
stat(s, 8.75, 1.9, 4.0, f"{br(vec_bw, 0)} GB/s", "vector: memória contígua, prefetch e SIMD", color=TEAL, vsize=40)
stat(s, 8.75, 3.5, 4.0, f"{br(list_bw)} GB/s", "list: cada nó é um ponteiro para outro lugar da memória",
     color=ORANGE, vsize=40)
box(s, 8.75, 5.3, 4.0, 1.6, TINT, radius=0.08)
text(s, 9.0, 5.45, 3.55, 1.4, [[("Lição: ", {"bold": True, "color": TEXT}),
                               ("a complexidade é igual; cache e layout de memória decidem o resultado.", {})]],
     size=13, color=MUTED)
notes(s, "SetBytesProcessed dá a vazão direto na saída. O vector passa de 30 GB/s; a list fica perto de 4, "
         "e piora com 1 milhão de elementos porque os nós deixam de caber no cache.")

# 16. Optimization: idea ----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Otimização guiada por benchmark", "Hipótese: a binária perde tempo com desvios mal previstos. Solução: tirar o if do laço.")
code(s, 0.6, 1.95, 6.0, 4.95, """
// antes: um if difícil de prever
while (lo < hi) {
    mid = lo + (hi - lo) / 2;
    if (values[mid] < target) lo = mid + 1;
    else                      hi = mid;
}

// depois: seleção sem desvio
while (len > 1) {
    half = len / 2;
    base = (base[half] < target)
               ? base + half : base;
    len -= half;
}
""", size=13)
code(s, 6.95, 1.95, 5.8, 2.25, """
$ objdump -d bin/Release/benchmarks
  cmp    (%rsi),%edi
  cmovg  %rsi,%rdx     <- sem jump
  sub    %rcx,%rax
""", size=13, label="CONFERINDO O ASSEMBLY", plain=True)
text(s, 6.95, 4.5, 5.8, 2.4, [
    "O laço executa sempre o mesmo número de passos; só o endereço muda.",
    "A CPU não precisa adivinhar o lado da busca, e cada erro de previsão custa da ordem de 15 ciclos.",
    "Os testes unitários confirmam que o resultado é igual ao da linear.",
], size=15, color=TEXT, spacing=12)
notes(s, "Com os números da pegadinha 2, dava para suspeitar dos desvios. A versão branchless troca o if por um "
         "cmov: o compilador gerou cmovg, conferido no objdump. Agora é medir se valeu a pena.")

# 17. Optimization: result --------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Resultado: compare.py", "Tempo por consulta (ns), mediana de 10 repetições")
chart(s, 0.45, 1.8, 7.4, 5.3, XL_CHART_TYPE.COLUMN_CLUSTERED, N_LABELS,
      [("Binária", D["binary"]), ("Branchless", D["branchless"])], [TEAL, ORANGE],
      number_format="0", labels=True, font_size=11)
speedup = sum(D["binary"]) / sum(D["branchless"])
geo = 1
for a, b in zip(D["binary"], D["branchless"]):
    geo *= b / a
geo = geo ** (1 / len(D["binary"]))
stat(s, 8.2, 1.85, 4.6, "−77%", "tempo (média geométrica em todos os N)", color=ORANGE, vsize=48)
stat(s, 8.2, 3.35, 4.6, "p = 0,0002", "Mann-Whitney U-test, 10 × 10 repetições", color=TEXT, vsize=32)
code(s, 8.2, 4.9, 4.55, 1.55, """
$ compare.py benchmarksfiltered \\
    v2.json BM_BinarySearch \\
    v2.json BM_BranchlessBinarySearch
OVERALL_GEOMEAN   -0.7748
""", size=10.5, plain=True)
notes(s, "O compare.py vem com a biblioteca. Ele compara as repetições das duas versões e roda um teste U de "
         "Mann-Whitney. O p-valor de 0,0002 é o mínimo possível com 10 contra 10 amostras: a diferença é real, "
         "cerca de 4,4 vezes mais rápido. A branchless em N=16 (3,7 ns) ganha até da linear.")

# 18. Checklist -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Checklist de boas práticas")
checks = [
    ("Release, flags de produção", "-O3, as mesmas flags do binário real"),
    ("Só o essencial no laço", "Setup fora; PauseTiming só se não houver alternativa"),
    ("DoNotOptimize no resultado", "E ClobberMemory quando escreve em memória"),
    ("Dados realistas e variados", "Padrões repetidos são aprendidos pela CPU"),
    ("Repetições + mediana + CV", "--benchmark_repetitions, CV abaixo de 2%"),
    ("Ambiente controlado", "taskset, governor performance, máquina ociosa"),
    ("Compare com estatística", "compare.py e p-valor, não comparar no olho"),
    ("Microbenchmark ≠ sistema", "Confirme no programa real com perf ou um profiler"),
]
for i, (h, b) in enumerate(checks):
    col, row = i % 2, i // 2
    x, yy = 0.6 + col * 6.2, 1.7 + row * 1.35
    badge(s, x, yy + 0.05, "✓", d=0.46, fill=TEAL, size=15)
    text(s, x + 0.65, yy, 5.3, 0.4, h, size=17, bold=True)
    text(s, x + 0.65, yy + 0.42, 5.3, 0.5, b, size=14, color=MUTED)
notes(s, "Resumo do que aprendi no piloto. O último ponto é importante: microbenchmark mede um trecho isolado, "
         "com cache quente. Serve para comparar alternativas, não para prever o tempo do sistema inteiro.")

# 19. Closing ---------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, INK)
text(s, 0.8, 1.9, 11, 1.2, "Obrigado!", size=54, bold=True, color=WHITE)
text(s, 0.8, 3.0, 11, 0.6, "Perguntas?", size=24, color="C9D1D9")
refs = [
    ("Repositório", "github.com/google/benchmark"),
    ("Guia do usuário", "github.com/google/benchmark/blob/main/docs/user_guide.md"),
    ("Comparação A/B", "github.com/google/benchmark/blob/main/docs/tools.md"),
    ("Código do piloto", "GBENCH_PILOT  (./run -B)"),
]
for i, (h, u) in enumerate(refs):
    yy = 4.3 + i * 0.55
    text(s, 0.8, yy, 2.6, 0.4, h, size=14, bold=True, color=ORANGE)
    text(s, 3.4, yy, 9.3, 0.4, u, size=14, color="D5DCE3", font=MONO)
notes(s, "Links para a documentação oficial. O guia do usuário cobre todos os recursos mostrados aqui.")

out = os.path.join(HERE, "google_benchmark.pptx")
prs.save(out)
print(out)
